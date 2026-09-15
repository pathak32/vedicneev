import json
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

SRC = sys.argv[1]
OUT = sys.argv[2]
BOARD_LABEL = sys.argv[3] if len(sys.argv) > 3 else ""

with open(SRC, encoding="utf-8") as f:
    rows = json.load(f)

wb = Workbook()
wb.remove(wb.active)

SECTION_ORDER = ["arithmetic", "mental_ability", "language", "general_knowledge"]
by_section = {}
for r in rows:
    by_section.setdefault(r["section"], []).append(r)

header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="C0392B", end_color="C0392B", fill_type="solid")
wrap = Alignment(wrap_text=True, vertical="top")
bn_fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
review_fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")

HEADERS = [
    ("Key", 22),
    ("Paper", 8),
    ("Difficulty", 10),
    ("Correct (0-3)", 12),
    ("Question (EN)", 40),
    ("Question (HI)", 32),
    ("Question (MR)", 32),
    ("Question (BN) - DRAFT", 40),
    ("Option A (EN)", 20), ("Option B (EN)", 20), ("Option C (EN)", 20), ("Option D (EN)", 20),
    ("Option A (BN) - DRAFT", 20), ("Option B (BN) - DRAFT", 20), ("Option C (BN) - DRAFT", 20), ("Option D (BN) - DRAFT", 20),
    ("Explanation (EN)", 45),
    ("Explanation (HI)", 32),
    ("Explanation (MR)", 32),
    ("Explanation (BN) - DRAFT", 45),
    ("Verified? (Y/N)", 14),
    ("Reviewer Notes", 35),
]

for section in SECTION_ORDER:
    if section not in by_section:
        continue
    sec_rows = by_section[section]
    ws = wb.create_sheet(title=section[:31])
    for col, (title, width) in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"

    for i, r in enumerate(sec_rows, start=2):
        opts_en = r["options_en"] + [""] * (4 - len(r["options_en"]))
        opts_bn = r["options_bn_draft"] + [""] * (4 - len(r["options_bn_draft"]))
        values = [
            r["key"], r["paperNumber"], r["difficulty"], r["correctAnswer"],
            r["question_en"], r["question_hi"], r["question_mr"], r["question_bn_draft"],
            opts_en[0], opts_en[1], opts_en[2], opts_en[3],
            opts_bn[0], opts_bn[1], opts_bn[2], opts_bn[3],
            r["explanation_en"], r["explanation_hi"], r["explanation_mr"], r["explanation_bn_draft"],
            "", "",
        ]
        for col, v in enumerate(values, start=1):
            cell = ws.cell(row=i, column=col, value=v)
            cell.alignment = wrap
            if col in (8, 13, 14, 15, 16, 20):
                cell.fill = bn_fill
            if col in (21, 22):
                cell.fill = review_fill

summary = wb.create_sheet(title="README", index=0)
summary["A1"] = f"Bengali (bn) PYQ translation review — {BOARD_LABEL}"
summary["A1"].font = Font(bold=True, size=14)
summary["A3"] = (
    "Each tab is one exam section. Columns in yellow are the DRAFT Bengali translation "
    "(machine-drafted by Claude, using the existing EN/HI/MR text as reference) — nothing here "
    "has been merged into pyq-seed/*.ts or the database yet."
)
summary["A3"].alignment = Alignment(wrap_text=True)
summary["A5"] = (
    "To review: read the EN column for meaning, compare against the BN draft, and fix any BN cell "
    "directly in this sheet. Mark 'Verified? (Y/N)' once a row is checked, and use 'Reviewer Notes' "
    "for anything that needs a second look before it goes live."
)
summary["A5"].alignment = Alignment(wrap_text=True)
summary.column_dimensions["A"].width = 110
for row in (1, 3, 5):
    summary.row_dimensions[row].height = 40 if row != 1 else 20

wb.save(OUT)
print("Wrote", OUT, "with sheets:", wb.sheetnames)
